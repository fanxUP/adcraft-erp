"""学校清单报价导入：清洗、计量识别、预览和批量写入准备。"""

from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.business_document import BusinessDocument, BusinessDocumentItem
from app.models.customer import Customer
from app.services.business_document_service import BusinessDocumentService


MONEY_QUANTUM = Decimal("0.01")
AREA_QUANTUM = Decimal("0.01")
DIMENSION_QUANTUM = Decimal("0.001")
PREVIEW_VERSION = "school-list-v1"
REQUIRED_HEADERS = (
    "部门/科室",
    "项目内容",
    "产品/材质/工艺",
    "数量",
    "单位",
    "单价",
)


class SchoolQuoteImportFormatError(ValueError):
    """清单表头或导入参数不符合学校清单模式。"""


_DIMENSION_NUMBER = r"[0-9]+(?:\.[0-9]+)?"
_DIMENSION_UNIT = r"(?:毫米|厘米|公分|mm|cm|米|m)"
_DIMENSION_PAIR_RE = re.compile(
    rf"(?P<first>{_DIMENSION_NUMBER})\s*(?P<first_unit>{_DIMENSION_UNIT})?\s*"
    rf"(?:×|x|X|\*)\s*"
    rf"(?P<second>{_DIMENSION_NUMBER})\s*(?P<second_unit>{_DIMENSION_UNIT})?"
    rf"(?!\s*(?:×|x|X|\*))",
    re.IGNORECASE,
)
_DIMENSION_TRIPLE_RE = re.compile(
    rf"(?P<first>{_DIMENSION_NUMBER})\s*(?P<first_unit>{_DIMENSION_UNIT})?\s*"
    rf"(?:×|x|X|\*)\s*"
    rf"(?P<second>{_DIMENSION_NUMBER})\s*(?P<second_unit>{_DIMENSION_UNIT})?\s*"
    rf"(?:×|x|X|\*)\s*"
    rf"(?P<third>{_DIMENSION_NUMBER})\s*(?P<third_unit>{_DIMENSION_UNIT})?",
    re.IGNORECASE,
)
_LABELLED_DIMENSION_RE = re.compile(
    rf"(?P<label>长|长度|宽|宽度|高|高度)\s*[:：]?\s*"
    rf"(?P<value>{_DIMENSION_NUMBER})\s*(?P<unit>{_DIMENSION_UNIT})?",
    re.IGNORECASE,
)
_DIAMETER_RE = re.compile(
    rf"(?:直径|Φ|φ)\s*[:：]?\s*(?P<value>{_DIMENSION_NUMBER})\s*"
    rf"(?P<unit>{_DIMENSION_UNIT})?",
    re.IGNORECASE,
)
_STRONG_DIMENSION_CONTEXT_RE = re.compile(
    r"尺寸|洞口尺寸|字体规格|箱体规格|台柜规格|门代号",
    re.IGNORECASE,
)
_GENERIC_DIMENSION_CONTEXT_RE = re.compile(r"规格", re.IGNORECASE)
_STRUCTURAL_CONTEXT_RE = re.compile(
    r"方钢|方管|骨架|主骨|副骨|间距|厚|基座|基础|埋深|地上|"
    r"面板选用|基材|立柱|格栅|龙骨|预埋",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class DimensionExtraction:
    """从产品/材质/工艺文本中提取的、可解释的结构化尺寸。"""

    status: str = "none"  # auto | review | none
    length: Decimal | None = None
    width: Decimal | None = None
    height: Decimal | None = None
    length_unit: str | None = None
    width_unit: str | None = None
    height_unit: str | None = None
    source: str | None = None
    reason: str | None = None


def _dimension_multiplier(unit: str | None) -> Decimal:
    normalized = (unit or "").strip().lower()
    if normalized in {"毫米", "mm"}:
        return Decimal("0.001")
    if normalized in {"厘米", "公分", "cm"}:
        return Decimal("0.01")
    return Decimal("1")


def _infer_dimension_unit(values: list[Decimal], explicit_units: list[str | None]) -> str:
    explicit = next((unit for unit in explicit_units if unit), None)
    if explicit:
        return explicit
    if any(value < Decimal("1") for value in values) or any(
        value != value.to_integral_value() for value in values
    ):
        return "m"
    if max(values, default=Decimal("0")) >= Decimal("100"):
        return "mm"
    return "m"


def _normalize_dimension(value: str, unit: str | None, inferred_unit: str) -> Decimal:
    number = Decimal(value)
    return (number * _dimension_multiplier(unit or inferred_unit)).quantize(
        DIMENSION_QUANTUM,
        rounding=ROUND_HALF_UP,
    )


def _line_can_contain_product_dimension(line: str) -> bool:
    """Reject material/profile dimensions unless the line names the product size."""
    if _DIAMETER_RE.search(line):
        return True
    if _STRONG_DIMENSION_CONTEXT_RE.search(line):
        return True
    if _STRUCTURAL_CONTEXT_RE.search(line):
        return False
    return bool(_GENERIC_DIMENSION_CONTEXT_RE.search(line))


def _dimension_from_line(line: str) -> DimensionExtraction | None:
    line = line.strip()
    if not line or not _line_can_contain_product_dimension(line):
        return None

    labelled = list(_LABELLED_DIMENSION_RE.finditer(line))
    if labelled:
        raw_values = [Decimal(match.group("value")) for match in labelled]
        inferred_unit = _infer_dimension_unit(
            raw_values,
            [match.group("unit") for match in labelled],
        )
        values = {
            match.group("label"): _normalize_dimension(
                match.group("value"),
                match.group("unit"),
                inferred_unit,
            )
            for match in labelled
        }
        length = values.get("长") or values.get("长度")
        width = values.get("宽") or values.get("宽度")
        height = values.get("高") or values.get("高度")
        if width is None and length is not None and height is not None:
            # 二维立面项目通常写作“长、高”，横向尺寸展示在“宽”列。
            width = length
            length = None
        if height is not None and width is not None:
            return DimensionExtraction(
                status="auto",
                length=length,
                width=width,
                height=height,
                length_unit="m" if length is not None else None,
                width_unit="m",
                height_unit="m",
                source=line,
                reason="识别到明确的长/宽/高标签",
            )

    triple = _DIMENSION_TRIPLE_RE.search(line)
    if triple and _STRONG_DIMENSION_CONTEXT_RE.search(line):
        raw_values = [
            Decimal(triple.group(name))
            for name in ("first", "second", "third")
        ]
        inferred_unit = _infer_dimension_unit(
            raw_values,
            [triple.group(f"{name}_unit") for name in ("first", "second", "third")],
        )
        return DimensionExtraction(
            status="review",
            length=_normalize_dimension(triple.group("first"), triple.group("first_unit"), inferred_unit),
            width=_normalize_dimension(triple.group("second"), triple.group("second_unit"), inferred_unit),
            height=_normalize_dimension(triple.group("third"), triple.group("third_unit"), inferred_unit),
            length_unit="m",
            width_unit="m",
            height_unit="m",
            source=triple.group(0),
            reason="三维尺寸需要确认长、宽、高的业务语义",
        )

    pair = _DIMENSION_PAIR_RE.search(line)
    if pair and _STRONG_DIMENSION_CONTEXT_RE.search(line):
        raw_values = [Decimal(pair.group("first")), Decimal(pair.group("second"))]
        inferred_unit = _infer_dimension_unit(
            raw_values,
            [pair.group("first_unit"), pair.group("second_unit")],
        )
        return DimensionExtraction(
            status="auto",
            width=_normalize_dimension(pair.group("first"), pair.group("first_unit"), inferred_unit),
            height=_normalize_dimension(pair.group("second"), pair.group("second_unit"), inferred_unit),
            width_unit="m",
            height_unit="m",
            source=pair.group(0),
            reason="识别到明确的二维产品尺寸",
        )

    diameter = _DIAMETER_RE.search(line)
    if diameter:
        inferred_unit = _infer_dimension_unit(
            [Decimal(diameter.group("value"))],
            [diameter.group("unit")],
        )
        value = _normalize_dimension(diameter.group("value"), diameter.group("unit"), inferred_unit)
        return DimensionExtraction(
            status="auto",
            width=value,
            height=value,
            width_unit="m",
            height_unit="m",
            source=diameter.group(0),
            reason="圆形直径映射为外接宽、高",
        )

    return None


def extract_product_dimensions(text: str) -> DimensionExtraction:
    """按产品语义提取尺寸，不把材料厚度或型材截面当作成品宽高。"""
    lines = [line.strip() for line in re.split(r"[\n；;]", text or "") if line.strip()]
    candidates = [candidate for line in lines if (candidate := _dimension_from_line(line))]
    if not candidates:
        return DimensionExtraction(
            status="none",
            reason="未找到可安全归属到成品宽高的尺寸",
        )

    # 同一行的重复描述只保留一个候选；不同尺寸含义冲突时不自动写入。
    unique = []
    seen = set()
    for candidate in candidates:
        key = (
            candidate.status,
            candidate.length,
            candidate.width,
            candidate.height,
            candidate.source,
        )
        if key not in seen:
            unique.append(candidate)
            seen.add(key)
    if len(unique) > 1:
        return DimensionExtraction(
            status="review",
            source="；".join(candidate.source or "" for candidate in unique),
            reason="同一明细存在多个可能的尺寸来源，需要确认归属",
        )
    return unique[0]


def _clean_scalar(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").replace("\u3000", " ")
    return " ".join(text.replace("\t", " ").split()).strip()


def _clean_detail_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    return text.strip()


def _parse_decimal(value: Any, *, row: int, field_name: str) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name}不是有效数字") from exc


def _normalize_unit(unit: str) -> tuple[str, str]:
    """Return (display unit, measure kind)."""
    normalized = unit.strip().lower().replace(" ", "")
    normalized = normalized.replace("²", "2").replace("³", "3")
    if normalized in {"m2", "㎡", "平方米", "平米"}:
        return "㎡", "area"
    if normalized in {"m3", "立方米", "立方"}:
        return unit, "volume"
    if normalized in {"m", "米"}:
        return unit, "length"
    if normalized in {"t", "吨"}:
        return unit, "weight"
    return unit, "quantity"


def _preview_id(file_bytes: bytes, customer_name: str, project_name: str) -> str:
    source_hash = sha256(file_bytes).hexdigest()
    identity = f"{PREVIEW_VERSION}:{source_hash}:{customer_name}:{project_name}"
    return sha256(identity.encode("utf-8")).hexdigest()


@dataclass
class SchoolQuoteImportItem:
    row_number: int
    item_name: str
    material_process: str
    quantity: Decimal
    source_unit: str
    unit: str
    measure_kind: str
    quantity_mode: str
    unit_price: Decimal
    remark: str | None
    subtotal_amount: Decimal
    length: Decimal | None = None
    width: Decimal | None = None
    height: Decimal | None = None
    length_unit: str | None = None
    width_unit: str | None = None
    height_unit: str | None = None
    dimension_status: str = "none"
    dimension_source: str | None = None
    dimension_reason: str | None = None

    def to_quote_item(self) -> dict[str, Any]:
        # Excel 的 m2 是已经给出的显式面积数量，不是宽×高几何面积。
        data = {
            "item_name": self.item_name,
            "material_process": self.material_process or None,
            "quantity": self.quantity,
            "unit": self.unit,
            "quantity_mode": self.quantity_mode,
            "use_area": False,
            "pieces": Decimal("1"),
            "unit_price": self.unit_price,
            "other_fee": Decimal("0"),
            "remark": self.remark,
            "sort_order": self.row_number,
        }
        # 复核项只在预览中展示候选值，不在无人确认时写入报价。
        if self.dimension_status == "auto":
            data.update({
                "length": self.length,
                "length_unit": self.length_unit,
                "width": self.width,
                "width_unit": self.width_unit,
                "height": self.height,
                "height_unit": self.height_unit,
            })
        return data

    def to_preview_dict(self) -> dict[str, Any]:
        return {
            "row": self.row_number,
            "item_name": self.item_name,
            "quantity": float(self.quantity),
            "source_unit": self.source_unit,
            "unit": self.unit,
            "measure_kind": self.measure_kind,
            "quantity_mode": self.quantity_mode,
            "unit_price": float(self.unit_price),
            "subtotal_amount": float(self.subtotal_amount),
            "remark": self.remark,
            "length": float(self.length) if self.length is not None else None,
            "length_unit": self.length_unit,
            "width": float(self.width) if self.width is not None else None,
            "width_unit": self.width_unit,
            "height": float(self.height) if self.height is not None else None,
            "height_unit": self.height_unit,
            "dimension_status": self.dimension_status,
            "dimension_source": self.dimension_source,
            "dimension_reason": self.dimension_reason,
        }


@dataclass
class SchoolQuoteImportGroup:
    department: str
    items: list[SchoolQuoteImportItem] = field(default_factory=list)

    @property
    def subtotal_amount(self) -> Decimal:
        return sum((item.subtotal_amount for item in self.items), Decimal("0")).quantize(
            MONEY_QUANTUM,
            rounding=ROUND_HALF_UP,
        )

    def to_preview_dict(self) -> dict[str, Any]:
        unit_counts: dict[str, int] = {}
        for item in self.items:
            unit_counts[item.unit] = unit_counts.get(item.unit, 0) + 1
        return {
            "department": self.department,
            "item_count": len(self.items),
            "area_item_count": sum(item.measure_kind == "area" for item in self.items),
            "dimension_auto_count": sum(item.dimension_status == "auto" for item in self.items),
            "dimension_review_count": sum(item.dimension_status == "review" for item in self.items),
            "subtotal_amount": float(self.subtotal_amount),
            "unit_counts": unit_counts,
            "items": [item.to_preview_dict() for item in self.items],
        }


@dataclass
class SchoolQuoteImportPreview:
    preview_id: str
    source_sha256: str
    customer_name: str
    project_name: str
    groups: list[SchoolQuoteImportGroup]
    skipped_rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)

    @property
    def school_count(self) -> int:
        return len(self.groups)

    @property
    def item_count(self) -> int:
        return sum(len(group.items) for group in self.groups)

    @property
    def total_amount(self) -> Decimal:
        return sum((group.subtotal_amount for group in self.groups), Decimal("0")).quantize(
            MONEY_QUANTUM,
            rounding=ROUND_HALF_UP,
        )

    @property
    def dimension_auto_count(self) -> int:
        return sum(
            item.dimension_status == "auto"
            for group in self.groups
            for item in group.items
        )

    @property
    def dimension_review_count(self) -> int:
        return sum(
            item.dimension_status == "review"
            for group in self.groups
            for item in group.items
        )

    @property
    def dimension_empty_count(self) -> int:
        return sum(
            item.dimension_status == "none"
            for group in self.groups
            for item in group.items
        )

    @property
    def valid(self) -> bool:
        return not self.errors and bool(self.groups)

    def to_response(self) -> dict[str, Any]:
        return {
            "preview_id": self.preview_id,
            "source_sha256": self.source_sha256,
            "customer_name": self.customer_name,
            "project_name": self.project_name,
            "valid": self.valid,
            "school_count": self.school_count,
            "item_count": self.item_count,
            "total_amount": float(self.total_amount),
            "dimension_auto_count": self.dimension_auto_count,
            "dimension_review_count": self.dimension_review_count,
            "dimension_empty_count": self.dimension_empty_count,
            "skipped_rows": self.skipped_rows,
            "errors": self.errors,
            "warnings": self.warnings,
            "schools": [group.to_preview_dict() for group in self.groups],
        }


def _merged_column_values(ws, column_index: int) -> dict[int, str]:
    values: dict[int, str] = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_col > column_index or merged.max_col < column_index:
            continue
        value = _clean_detail_text(ws.cell(merged.min_row, column_index).value)
        if not value:
            continue
        for row_number in range(merged.min_row, merged.max_row + 1):
            values[row_number] = value
    return values


def parse_school_quote_workbook(
    file_bytes: bytes,
    *,
    customer_name: str,
    project_name: str,
) -> SchoolQuoteImportPreview:
    customer_name = _clean_scalar(customer_name)
    project_name = _clean_scalar(project_name)
    if not customer_name:
        raise SchoolQuoteImportFormatError("客户名称不能为空")
    if not project_name:
        raise SchoolQuoteImportFormatError("项目名称不能为空")

    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception as exc:
        raise SchoolQuoteImportFormatError("无法读取 Excel 文件") from exc

    try:
        worksheet = workbook.active
        if worksheet is None:
            raise SchoolQuoteImportFormatError("Excel 文件为空")
        headers = [_clean_scalar(cell.value) for cell in worksheet[1]]
        header_index = {header: index + 1 for index, header in enumerate(headers) if header}
        missing = [header for header in REQUIRED_HEADERS if header not in header_index]
        if missing:
            raise SchoolQuoteImportFormatError(f"缺少必填列: {', '.join(missing)}")

        remark_column = header_index.get("备注")
        merged_remarks = _merged_column_values(worksheet, remark_column) if remark_column else {}
        subtotal_column = header_index.get("小计")
        groups: OrderedDict[str, SchoolQuoteImportGroup] = OrderedDict()
        skipped_rows: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []

        for row_number in range(2, worksheet.max_row + 1):
            def cell(header: str) -> Any:
                return worksheet.cell(row_number, header_index[header]).value

            department = _clean_scalar(cell("部门/科室"))
            item_name = _clean_scalar(cell("项目内容"))
            if not department and not item_name and all(
                worksheet.cell(row_number, column).value is None
                for column in range(1, worksheet.max_column + 1)
            ):
                skipped_rows.append({"row": row_number, "reason": "空行"})
                continue
            if item_name == "合计" and not department:
                skipped_rows.append({"row": row_number, "reason": "汇总行"})
                continue

            row_errors: list[str] = []
            if not department:
                row_errors.append("部门/科室不能为空")
            if not item_name:
                row_errors.append("项目内容不能为空")
            raw_unit = _clean_scalar(cell("单位"))
            if not raw_unit:
                row_errors.append("单位不能为空")

            try:
                quantity = _parse_decimal(cell("数量"), row=row_number, field_name="数量")
            except ValueError as exc:
                quantity = None
                row_errors.append(str(exc))
            if quantity is None:
                row_errors.append("数量不能为空")
            elif quantity <= 0:
                row_errors.append("数量必须大于 0")

            try:
                unit_price = _parse_decimal(cell("单价"), row=row_number, field_name="单价")
            except ValueError as exc:
                unit_price = None
                row_errors.append(str(exc))
            if unit_price is None:
                row_errors.append("单价不能为空")
            elif unit_price < 0:
                row_errors.append("单价不能为负数")

            if row_errors:
                errors.append({"row": row_number, "message": "；".join(dict.fromkeys(row_errors))})
                continue

            display_unit, measure_kind = _normalize_unit(raw_unit)
            quantity_mode = "area" if measure_kind == "area" else "piece"
            subtotal_amount = (quantity * unit_price).quantize(
                MONEY_QUANTUM,
                rounding=ROUND_HALF_UP,
            )
            if subtotal_column:
                try:
                    source_subtotal = _parse_decimal(
                        worksheet.cell(row_number, subtotal_column).value,
                        row=row_number,
                        field_name="小计",
                    )
                except ValueError as exc:
                    warnings.append({"row": row_number, "message": str(exc)})
                else:
                    if source_subtotal is not None:
                        source_subtotal = source_subtotal.quantize(
                            MONEY_QUANTUM,
                            rounding=ROUND_HALF_UP,
                        )
                        if source_subtotal != subtotal_amount:
                            warnings.append({
                                "row": row_number,
                                "message": f"原始小计 {source_subtotal} 与数量×单价 {subtotal_amount} 不一致，已按数量×单价重算",
                            })

            remark = _clean_detail_text(cell("备注")) if remark_column else ""
            if not remark and row_number in merged_remarks:
                remark = merged_remarks[row_number]
            dimensions = extract_product_dimensions(_clean_detail_text(cell("产品/材质/工艺")))
            item = SchoolQuoteImportItem(
                row_number=row_number,
                item_name=item_name,
                material_process=_clean_detail_text(cell("产品/材质/工艺")),
                quantity=quantity,
                source_unit=raw_unit,
                unit=display_unit,
                measure_kind=measure_kind,
                quantity_mode=quantity_mode,
                unit_price=unit_price,
                remark=remark or None,
                subtotal_amount=subtotal_amount,
                length=dimensions.length,
                width=dimensions.width,
                height=dimensions.height,
                length_unit=dimensions.length_unit,
                width_unit=dimensions.width_unit,
                height_unit=dimensions.height_unit,
                dimension_status=dimensions.status,
                dimension_source=dimensions.source,
                dimension_reason=dimensions.reason,
            )
            groups.setdefault(department, SchoolQuoteImportGroup(department=department)).items.append(item)

        if not groups and not errors:
            errors.append({"row": 0, "message": "未找到有效学校明细"})

        source_sha256 = sha256(file_bytes).hexdigest()
        return SchoolQuoteImportPreview(
            preview_id=_preview_id(file_bytes, customer_name, project_name),
            source_sha256=source_sha256,
            customer_name=customer_name,
            project_name=project_name,
            groups=list(groups.values()),
            skipped_rows=skipped_rows,
            errors=errors,
            warnings=warnings,
        )
    finally:
        workbook.close()


async def resolve_customer_id(db: AsyncSession, customer_name: str):
    result = await db.execute(
        select(Customer.id)
        .where(Customer.name == customer_name, Customer.deleted_at.is_(None))
        .order_by(Customer.created_at.asc())
    )
    customer_ids = list(result.scalars().all())
    if len(customer_ids) > 1:
        raise ValueError("客户名称匹配到多个客户档案，请先合并客户档案后再导入")
    return customer_ids[0] if customer_ids else None


async def find_existing_school_quotes(
    db: AsyncSession,
    *,
    customer_name: str,
    customer_id,
    project_name: str,
    departments: list[str],
) -> list[dict[str, Any]]:
    customer_match = BusinessDocument.customer_name == customer_name
    if customer_id:
        customer_match = or_(
            BusinessDocument.customer_id == customer_id,
            BusinessDocument.customer_name == customer_name,
        )
    result = await db.execute(
        select(BusinessDocument)
        .where(
            BusinessDocument.doc_type == "quote",
            BusinessDocument.deleted_at.is_(None),
            BusinessDocument.project_name == project_name,
            BusinessDocument.department.in_(departments),
            customer_match,
        )
        .order_by(BusinessDocument.department, BusinessDocument.created_at),
    )
    return [
        {
            "id": str(quote.id),
            "quote_no": quote.doc_no,
            "department": quote.department,
            "total_amount": float(quote.total_amount or 0),
        }
        for quote in result.scalars().all()
    ]


def _normalized_match_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").replace("\u3000", " ").split())


def _same_decimal(left: Any, right: Any) -> bool:
    if left is None or right is None:
        return left is None and right is None
    return Decimal(str(left)) == Decimal(str(right))


def _same_dimension_value(left: Any, right: Any) -> bool:
    if left is None or right is None:
        return left is None and right is None
    return Decimal(str(left)).quantize(DIMENSION_QUANTUM) == Decimal(str(right)).quantize(DIMENSION_QUANTUM)


def _school_dimension_item_mismatch(source: SchoolQuoteImportItem, target: BusinessDocumentItem) -> list[str]:
    mismatches: list[str] = []
    if _normalized_match_text(source.item_name) != _normalized_match_text(target.item_name):
        mismatches.append("项目内容")
    if not _same_decimal(source.quantity, target.quantity):
        mismatches.append("数量")
    if _normalized_match_text(source.unit) != _normalized_match_text(target.unit):
        mismatches.append("单位")
    if not _same_decimal(source.unit_price, target.unit_price):
        mismatches.append("单价")
    if _normalized_match_text(source.material_process) != _normalized_match_text(target.material_process):
        mismatches.append("产品/材质/工艺")
    return mismatches


def _dimension_backfill_fields(source: SchoolQuoteImportItem) -> dict[str, Any]:
    return {
        "length": source.length,
        "length_unit": source.length_unit,
        "width": source.width,
        "width_unit": source.width_unit,
        "height": source.height,
        "height_unit": source.height_unit,
    }


def _dimension_preview_fields(source: SchoolQuoteImportItem) -> dict[str, Any]:
    values = _dimension_backfill_fields(source)
    return {
        key: float(value) if key in {"length", "width", "height"} and value is not None else value
        for key, value in values.items()
    }


async def _load_school_quotes_for_dimension_backfill(
    db: AsyncSession,
    preview: SchoolQuoteImportPreview,
    *,
    customer_id=None,
) -> list[BusinessDocument]:
    customer_match = BusinessDocument.customer_name == preview.customer_name
    if customer_id:
        customer_match = or_(
            BusinessDocument.customer_id == customer_id,
            BusinessDocument.customer_name == preview.customer_name,
        )
    result = await db.execute(
        select(BusinessDocument)
        .where(
            BusinessDocument.doc_type == "quote",
            BusinessDocument.deleted_at.is_(None),
            BusinessDocument.project_name == preview.project_name,
            BusinessDocument.department.in_([group.department for group in preview.groups]),
            customer_match,
        )
        .order_by(BusinessDocument.department, BusinessDocument.created_at),
    )
    return list(result.scalars().all())


@dataclass
class _DimensionBackfillPlan:
    response: dict[str, Any]
    changes: list[tuple[BusinessDocumentItem, dict[str, Any]]] = field(default_factory=list)


async def _build_school_dimension_backfill_plan(
    db: AsyncSession,
    preview: SchoolQuoteImportPreview,
    *,
    customer_id=None,
) -> _DimensionBackfillPlan:
    quotes = await _load_school_quotes_for_dimension_backfill(
        db,
        preview,
        customer_id=customer_id,
    )
    quotes_by_department: dict[str, list[BusinessDocument]] = {}
    for quote in quotes:
        quotes_by_department.setdefault(quote.department or "", []).append(quote)

    changes: list[tuple[BusinessDocumentItem, dict[str, Any]]] = []
    change_rows: list[dict[str, Any]] = []
    review_rows: list[dict[str, Any]] = []
    skipped_rows: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    matched_quote_count = 0
    unchanged_count = 0

    for group in preview.groups:
        matches = quotes_by_department.get(group.department, [])
        if len(matches) != 1:
            conflicts.append({
                "department": group.department,
                "reason": "未找到唯一的目标草稿报价" if not matches else "同一学校匹配到多张报价",
                "quote_count": len(matches),
            })
            continue

        quote = matches[0]
        if quote.status != "draft":
            conflicts.append({
                "department": group.department,
                "quote_no": quote.doc_no,
                "reason": f"报价状态为 {quote.status}，只允许回填草稿报价",
            })
            continue
        matched_quote_count += 1

        source_items = list(group.items)
        target_items = sorted(
            list(quote.items),
            key=lambda item: (getattr(item, "sort_order", 0) or 0, str(getattr(item, "id", ""))),
        )
        if len(source_items) != len(target_items):
            conflicts.append({
                "department": group.department,
                "quote_no": quote.doc_no,
                "reason": "报价明细数量与原始清单不一致",
                "source_item_count": len(source_items),
                "target_item_count": len(target_items),
            })
            continue

        quote_has_conflict = False
        quote_pairs: list[tuple[SchoolQuoteImportItem, BusinessDocumentItem]] = []
        for source_item, target_item in zip(source_items, target_items):
            mismatches = _school_dimension_item_mismatch(source_item, target_item)
            if mismatches:
                quote_has_conflict = True
                conflicts.append({
                    "department": group.department,
                    "quote_no": quote.doc_no,
                    "row": source_item.row_number,
                    "item_name": source_item.item_name,
                    "reason": "报价明细与原始清单不一致",
                    "fields": mismatches,
                })
            else:
                quote_pairs.append((source_item, target_item))
        if quote_has_conflict:
            continue

        for source_item, target_item in quote_pairs:
            base = {
                "department": group.department,
                "quote_no": quote.doc_no,
                "item_id": str(target_item.id),
                "row": source_item.row_number,
                "item_name": source_item.item_name,
                "dimension_source": source_item.dimension_source,
            }
            if source_item.dimension_status == "review":
                review_rows.append({
                    **base,
                    "reason": source_item.dimension_reason,
                    **_dimension_preview_fields(source_item),
                })
                continue
            if source_item.dimension_status != "auto":
                skipped_rows.append({
                    **base,
                    "reason": source_item.dimension_reason or "没有可安全回填的尺寸",
                })
                continue

            planned_fields: dict[str, Any] = {}
            for field_name, proposed in _dimension_backfill_fields(source_item).items():
                if proposed is None:
                    continue
                current = getattr(target_item, field_name, None)
                is_dimension = field_name in {"length", "width", "height"}
                same = _same_dimension_value(current, proposed) if is_dimension else _normalized_match_text(current) == _normalized_match_text(proposed)
                if current is None:
                    planned_fields[field_name] = proposed
                elif not same:
                    conflicts.append({
                        **base,
                        "reason": f"目标明细已有不同的{field_name}，不会覆盖人工数据",
                        "current": str(current),
                        "proposed": str(proposed),
                    })
            if any(conflict.get("item_id") == str(target_item.id) for conflict in conflicts):
                continue
            if planned_fields:
                changes.append((target_item, planned_fields))
                change_rows.append({
                    **base,
                    **_dimension_preview_fields(source_item),
                    "fields": {key: str(value) for key, value in planned_fields.items()},
                })
            else:
                unchanged_count += 1

    response = {
        "valid": not conflicts and matched_quote_count == len(preview.groups),
        "preview_id": preview.preview_id,
        "customer_name": preview.customer_name,
        "project_name": preview.project_name,
        "expected_quote_count": len(preview.groups),
        "matched_quote_count": matched_quote_count,
        "change_count": len(changes),
        "unchanged_count": unchanged_count,
        "review_count": len(review_rows),
        "skipped_count": len(skipped_rows),
        "conflict_count": len(conflicts),
        "changes": change_rows,
        "review_rows": review_rows,
        "skipped_rows": skipped_rows,
        "conflicts": conflicts,
    }
    return _DimensionBackfillPlan(response=response, changes=changes)


async def preview_school_quote_dimension_backfill(
    db: AsyncSession,
    preview: SchoolQuoteImportPreview,
    *,
    customer_id=None,
) -> dict[str, Any]:
    """只读生成现有学校报价的尺寸回填计划。"""
    return (
        await _build_school_dimension_backfill_plan(
            db,
            preview,
            customer_id=customer_id,
        )
    ).response


async def commit_school_quote_dimension_backfill(
    db: AsyncSession,
    preview: SchoolQuoteImportPreview,
    *,
    customer_id=None,
) -> dict[str, Any]:
    """在匹配、状态和字段一致性全部通过后，事务内回填结构化尺寸。"""
    plan = await _build_school_dimension_backfill_plan(
        db,
        preview,
        customer_id=customer_id,
    )
    if not plan.response["valid"]:
        raise ValueError("尺寸回填预览存在匹配冲突，未写入任何数据")
    for item, fields in plan.changes:
        for field_name, value in fields.items():
            setattr(item, field_name, value)
    if plan.changes:
        await db.flush()
    return plan.response


async def commit_school_quote_preview(
    db: AsyncSession,
    preview: SchoolQuoteImportPreview,
    *,
    customer_id=None,
) -> list[dict[str, Any]]:
    if not preview.valid:
        raise ValueError("清单预览存在错误，不能提交")
    service = BusinessDocumentService(db, doc_type="quote")
    created: list[dict[str, Any]] = []
    for group in preview.groups:
        quote_data: dict[str, Any] = {
            "customer_id": str(customer_id) if customer_id else None,
            "customer_name": None if customer_id else preview.customer_name,
            "project_name": preview.project_name,
            "department": group.department,
            "discount_amount": Decimal("0"),
            "tax_rate": Decimal("0"),
            "items": [item.to_quote_item() for item in group.items],
        }
        quote = await service.create(quote_data)
        created.append({
            "id": str(quote["id"]),
            "quote_no": quote["quote_no"],
            "department": group.department,
            "item_count": len(group.items),
            "total_amount": float(quote["total_amount"]),
        })
    return created
